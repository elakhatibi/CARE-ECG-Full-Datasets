import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import numpy as np

from care_ecg.cli import (
    SECURITY_CONFIG_SCHEMA,
    SECURITY_RESULTS_SCHEMA,
    _directory_identity,
    _f0_magnitude_calibration,
    _f0_magnitudes_for_constraint,
    _load_security_config,
    _security_attack_factory,
    _security_code_identity,
    _security_cohort_cases,
    _security_lock_payload,
    _security_paired_f3_f1,
    _security_paired_f3_f1_automated_nonclinical,
    _security_raw_signals,
    _validate_security_lock,
    _write_security_workbook,
    build_parser,
)
from care_ecg.data import ECGArrayBundle


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SecurityCLITests(unittest.TestCase):
    def test_security_config_and_parser_expose_auditable_commands(self):
        config = _load_security_config("config/security_ptbxl.yaml")
        self.assertEqual(config["schema_version"], SECURITY_CONFIG_SCHEMA)
        self.assertGreaterEqual(
            set(config["attacks"]["enabled"]), {"f0", "f1_pgd", "f2_nes", "f3_k3"}
        )

        args = build_parser().parse_args(
            [
                "security-evaluate",
                "artifacts/model",
                "data/cohort.npz",
                "results/security",
                "--attacks",
                "f0,f3_k3",
                "--constraints",
                "s_inf,s_phys",
                "--budgets",
                "1,10",
                "--no-excel",
            ]
        )
        self.assertEqual(args.attacks, "f0,f3_k3")
        self.assertEqual(args.constraints, "s_inf,s_phys")
        self.assertFalse(args.excel)

    def test_security_cohort_maps_targets_and_deduplicates_patients(self):
        signals = np.zeros((4, 1, 8), dtype=np.float32)
        bundle = ECGArrayBundle(
            signals=signals,
            labels=np.asarray([2, 3, 3, 0]),
            patient_ids=np.asarray(["p1", "p2", "p2", "p3"]),
            record_ids=np.asarray(["r1", "r2", "r3", "r4"]),
            folds=np.asarray([9, 9, 9, 9]),
        )
        cases = _security_cohort_cases(
            bundle,
            signals,
            ("CD", "HYP", "MI", "NORM", "STTC"),
            ("T1", "T3", "T4"),
        )

        self.assertEqual(
            [(case.target.target_id, case.record_id) for case in cases],
            [("T1", "r1"), ("T3", "r2"), ("T4", "r2")],
        )
        self.assertTrue(all(case.metadata["fold"] == 9 for case in cases))

    def test_security_cohort_requires_patient_ids_for_primary_analysis(self):
        signals = np.zeros((1, 1, 8), dtype=np.float32)
        bundle = ECGArrayBundle(
            signals=signals,
            labels=np.asarray([2]),
            record_ids=np.asarray(["r1"]),
        )
        with self.assertRaisesRegex(ValueError, "Patient IDs"):
            _security_cohort_cases(
                bundle,
                signals,
                ("CD", "HYP", "MI", "NORM", "STTC"),
                ("T1",),
            )

        cases = _security_cohort_cases(
            bundle,
            signals,
            ("CD", "HYP", "MI", "NORM", "STTC"),
            ("T1",),
            allow_record_level=True,
        )
        self.assertEqual(cases[0].patient_id, "r1")

    def test_security_raw_signal_preparation_preserves_raw_mv(self):
        pipeline = SimpleNamespace(
            config=SimpleNamespace(
                data=SimpleNamespace(leads=1, samples=4, sample_rate=2, seconds=2)
            )
        )
        signals = np.asarray([[[0.0, 1.0, 2.0, 3.0]]], dtype=np.float32)
        prepared = _security_raw_signals(signals, pipeline, source_rate=None)
        np.testing.assert_array_equal(prepared, signals)

        with self.assertRaisesRegex(ValueError, "provide --source-rate"):
            _security_raw_signals(np.zeros((1, 1, 3), dtype=np.float32), pipeline, None)

    def test_security_report_computes_paired_f3_minus_f1(self):
        def run(successes):
            return {
                "observations": [
                    {
                        "patient_id": patient,
                        "record_id": f"r{index}",
                        "target_id": "T1",
                        "seed": 42,
                        "budget": 10,
                        "defense": "v1",
                        "success": success,
                        "metrics": {"primary_endpoint_eligible": True},
                    }
                    for index, (patient, success) in enumerate(
                        zip(("p1", "p2"), successes, strict=True)
                    )
                ]
            }

        paired = _security_paired_f3_f1(
            {
                "f1_pgd": {"s_inf": run((False, False))},
                "f3_k3": {"s_inf": run((True, False))},
            }
        )
        estimate = paired["f3_k3-f1_pgd:s_inf|T=T1|B=10|v1"]
        self.assertEqual(estimate["estimate"], 0.5)
        self.assertEqual(estimate["metadata"]["direction"], "adaptive_minus_comparator")
        self.assertEqual(estimate["metadata"]["target_id"], "T1")
        self.assertEqual(estimate["metadata"]["paired_observations"], 2)

    def test_security_report_rejects_incomplete_f3_f1_pairs(self):
        baseline = {
            "observations": [
                {
                    "patient_id": "p1",
                    "record_id": "r1",
                    "target_id": "T1",
                    "seed": 42,
                    "budget": 10,
                    "defense": "v1",
                    "success": False,
                }
            ]
        }

        with self.assertRaisesRegex(ValueError, "Incomplete .* pairing"):
            _security_paired_f3_f1(
                {
                    "f1_pgd": {"s_inf": baseline},
                    "f3_k3": {"s_inf": {"observations": []}},
                }
            )

    def test_security_report_computes_paired_automated_nonclinical_deltas(self):
        def run(values):
            observations = []
            for index, (
                patient,
                seed,
                targeted,
                released_claim,
                robust,
                clean_correct,
            ) in enumerate(values):
                observations.append(
                    {
                        "patient_id": patient,
                        "record_id": f"r{patient}",
                        "target_id": "T1",
                        "seed": seed,
                        "budget": 10,
                        "defense": "v1",
                        "success": False,
                        "metrics": {
                            "primary_endpoint_eligible": False,
                            "automated_robustness": {
                                "clean_correct": clean_correct,
                                "targeted_diagnosis_success": targeted,
                                "released_target_claim_success": released_claim,
                                "robust_correct": robust,
                            },
                        },
                        "row": index,
                    }
                )
            return {"observations": observations}

        comparator = run(
            (
                ("p1", 1, False, False, True, True),
                ("p1", 2, False, False, True, True),
                ("p2", 1, False, False, True, True),
                ("p3", 1, False, False, True, False),
            )
        )
        adaptive = run(
            (
                ("p1", 1, True, True, False, True),
                ("p1", 2, True, False, False, True),
                ("p2", 1, False, False, True, True),
                ("p3", 1, True, True, True, False),
            )
        )
        paired = _security_paired_f3_f1_automated_nonclinical(
            {
                "f1_pgd": {"s_inf": comparator},
                "f3_k3": {"s_inf": adaptive},
            }
        )

        self.assertEqual(paired["endpoint_classification"], "automated_nonclinical")
        self.assertTrue(paired["clinical_primary_endpoint_unchanged"])
        self.assertEqual(paired["multiplicity"]["adjustment"], "none")
        self.assertEqual(paired["multiplicity"]["planned_estimates"], 3)
        self.assertEqual(paired["multiplicity"]["evaluable_estimates"], 3)
        stratum = paired["strata"]["f3_k3-f1_pgd:s_inf|T=T1|B=10|v1"]
        targeted = stratum["metrics"]["targeted_diagnosis_asr_delta"]
        released = stratum["metrics"]["released_target_claim_asr_delta"]
        robust = stratum["metrics"]["robust_accuracy_delta"]
        self.assertEqual(targeted["estimate"], 0.5)
        self.assertEqual(released["estimate"], 0.25)
        self.assertAlmostEqual(robust["estimate"], -1.0 / 3.0)
        self.assertEqual(targeted["n_patients"], 2)
        self.assertEqual(targeted["metadata"]["paired_observations"], 3)
        self.assertEqual(robust["metadata"]["paired_observations"], 4)
        self.assertEqual(robust["n_patients"], 3)
        self.assertEqual(targeted["metadata"]["adaptive_events"], 2)
        self.assertEqual(targeted["metadata"]["comparator_events"], 0)
        self.assertEqual(
            targeted["metadata"]["conditioning"],
            "joint_clean_correct_in_adaptive_and_comparator",
        )
        self.assertEqual(
            robust["metadata"]["conditioning"],
            "all_matched_pairs_with_automated_endpoint",
        )

    def test_security_workbook_exports_paired_automated_deltas(self):
        estimate = {
            "estimate": 0.25,
            "lower": 0.0,
            "upper": 0.5,
            "confidence": 0.95,
            "n_patients": 2,
            "metadata": {
                "direction": "adaptive_minus_comparator",
                "paired_observations": 3,
                "adaptive_events": 1,
                "comparator_events": 0,
                "conditioning": "joint_clean_correct_in_adaptive_and_comparator",
                "denominator_unit": "matched observation pairs",
                "endpoint_classification": "automated_nonclinical",
                "multiplicity_adjustment": "none",
                "confidence_interval_scope": "nominal pointwise",
            },
        }
        paired = {
            "multiplicity": {
                "planned_estimates": 3,
                "evaluable_estimates": 3,
                "interpretation": "exploratory; no multiplicity control",
            },
            "strata": {
                "stratum": {
                    "adaptive_attack": "f3_k3",
                    "comparator_attack": "f1_pgd",
                    "constraint": "s_inf",
                    "target_id": "T1",
                    "budget": 10,
                    "defense": "v1",
                    "metrics": {"targeted_diagnosis_asr_delta": estimate},
                }
            },
        }
        report = {
            "e0": {"summary": {}, "observations": []},
            "e1": {},
            "paired_f3_minus_f1": {},
            "paired_f3_minus_f1_automated_nonclinical": paired,
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "security.xlsx"
            _write_security_workbook(report, path)
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True)
            sheet = workbook["Paired automated deltas"]
            rows = sheet.iter_rows()
            header = [cell.value for cell in next(rows)]
            row = [cell.value for cell in next(rows)]
            workbook.close()

        self.assertEqual(row[header.index("metric")], "targeted_diagnosis_asr_delta")
        self.assertEqual(row[header.index("paired_observations")], 3)
        self.assertEqual(row[header.index("endpoint_classification")], "automated_nonclinical")
        self.assertEqual(row[header.index("multiplicity_adjustment")], "none")
        self.assertEqual(row[header.index("multiplicity_planned_estimates")], 3)
        self.assertEqual(row[header.index("multiplicity_evaluable_estimates")], 3)
        self.assertEqual(
            row[header.index("multiplicity_interpretation")],
            "exploratory; no multiplicity control",
        )

    def test_f0_calibration_freezes_native_validation_median(self):
        def observation(distance):
            return {
                "budget": 10,
                "attack_result": {
                    "constraint_distances": [distance],
                    "constraint_satisfied": [True],
                },
            }

        calibration = _f0_magnitude_calibration(
            {
                "f1_pgd": {"s_inf": {"observations": [observation(0.01)]}},
                "f3_k3": {"s_inf": {"observations": [observation(0.03)]}},
            },
            constraints=("s_inf",),
            budgets=(10,),
            source_attacks=("f1_pgd", "f3_k3"),
        )

        self.assertEqual(_f0_magnitudes_for_constraint(calibration, "s_inf", (10,)), {10: 0.02})
        self.assertEqual(calibration["strata"]["s_inf|B=10"]["sample_count"], 2)
        calibration["strata"]["s_inf|B=10"]["magnitude"] = 0.5
        with self.assertRaisesRegex(ValueError, "hash mismatch"):
            _f0_magnitudes_for_constraint(calibration, "s_inf", (10,))

    def test_validation_calibration_overrides_direct_f0_magnitude(self):
        from care_ecg.security.constraints import SInfConstraint

        _, factory = _security_attack_factory(
            "f0",
            {"common": {}, "f0": {"matched_magnitude": 0.9}},
            SInfConstraint(0.02),
            matched_magnitudes={1: 0.01},
        )

        self.assertEqual(factory(seed=42, budget=1).matched_magnitude, 0.01)

    def test_gradient_attack_factory_propagates_momentum_decay(self):
        from care_ecg.security.constraints import SInfConstraint

        _, factory = _security_attack_factory(
            "f1_pgd",
            {"common": {"momentum_decay": 0.9}},
            SInfConstraint(0.02),
        )

        self.assertEqual(factory(seed=42, budget=1).momentum_decay, 0.9)

    def test_validation_lock_binds_config_and_validation_result(self):
        from care_ecg.security.logging import AppendOnlySecurityLogger
        from care_ecg.security.preferences import PreferencePolicy

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            dataset_path = root / "validation.npz"
            dataset_path.write_bytes(b"locked validation dataset")
            artifact_path = root / "artifacts"
            artifact_path.mkdir()
            (artifact_path / "model.bin").write_bytes(b"locked victim")
            config_path = root / "security.yaml"
            policy_path = root / "verifier.json"
            policy = PreferencePolicy("verifier", ["x"], ["release", "reject"], ["clean", "target"])
            policy.save(policy_path)
            config_path.write_text(
                "schema_version: care-ecg-security-config.v1\n"
                "protocol:\n"
                "  validation_folds: [9]\n"
                f"  verifier_policy: {policy_path}\n"
                f"  validation_dataset_sha256: {_sha256(dataset_path)}\n"
                "attacks: {}\n"
                "constraints: {}\nadmissibility: {}\nexport: {}\n",
                encoding="utf-8",
            )
            log_path = root / "security_events.jsonl"
            logger = AppendOnlySecurityLogger(log_path)
            logger.append("e0_clean", {"case": "p1:r1:T1"})
            logger.seal()
            verification = logger.verify(require_manifest=True).to_dict()
            validation_path = root / "validation.json"
            validation_path.write_text(
                json.dumps(
                    {
                        "schema_version": SECURITY_RESULTS_SCHEMA,
                        "protocol": {
                            "phase": "validation",
                            "security_config_sha256": _sha256(config_path),
                            "observed_folds": [9],
                            "candidate_cases": 1,
                            "eligible_cases_used": 1,
                            "attacks": [],
                            "constraints": [],
                            "preference_policies": {
                                "verifier": {
                                    "path": str(policy_path.resolve()),
                                    "file_sha256": _sha256(policy_path),
                                    "policy_sha256": policy.fingerprint(),
                                }
                            },
                        },
                        "artifacts": {
                            "path": str(artifact_path),
                            "identity": _directory_identity(artifact_path),
                        },
                        "implementation": {"code_identity": _security_code_identity()},
                        "dataset": {
                            "path": str(dataset_path),
                            "sha256": _sha256(dataset_path),
                            "preprocessing": {
                                "input_units": "mV",
                                "source_rate_hz": 100,
                            },
                        },
                        "event_log": {
                            "path": str(log_path),
                            "sha256": _sha256(log_path),
                            "manifest_path": str(logger.manifest_path),
                            "manifest_sha256": _sha256(logger.manifest_path),
                            "verification": verification,
                        },
                        "e0": {"observations": [{"patient_id": "p1"}]},
                        "e1": {},
                    }
                ),
                encoding="utf-8",
            )
            payload = _security_lock_payload(config_path, validation_path)
            lock_path = root / "lock.json"
            lock_path.write_text(json.dumps(payload), encoding="utf-8")

            verified = _validate_security_lock(lock_path, config_path)
            self.assertEqual(verified["validation_results_sha256"], _sha256(validation_path))
            self.assertEqual(
                verified["resolved_protocol"]["preference_policies"]["verifier"]["file_sha256"],
                _sha256(policy_path),
            )
            loaded_report = json.loads(validation_path.read_text())
            binding = loaded_report["protocol"].pop("preference_policies")
            validation_path.write_text(json.dumps(loaded_report))
            with self.assertRaisesRegex(ValueError, "bind all configured preference"):
                _security_lock_payload(config_path, validation_path)
            loaded_report["protocol"]["preference_policies"] = binding
            validation_path.write_text(json.dumps(loaded_report))
            policy.save(policy_path, metadata={"changed": True})
            with self.assertRaisesRegex(ValueError, "Preference policy changed"):
                _security_lock_payload(config_path, validation_path)

            config_path.write_text(config_path.read_text() + "# changed\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "changed after validation"):
                _validate_security_lock(lock_path, config_path)


if __name__ == "__main__":
    unittest.main()
